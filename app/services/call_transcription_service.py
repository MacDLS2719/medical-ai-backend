import io
import logging
import os
import uuid

import cloudinary
import httpx
from fastapi import UploadFile
from openai import AsyncOpenAI
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.medical_conversation import MedicalConversation
from app.services.deepgram_service import transcribe_audio
from app.services.medical_conversation_service import MedicalConversationService

logger = logging.getLogger(__name__)


class CallTranscriptionPipeline:
    """
    Orquesta el flujo completo de videollamada:
    Daily room -> audio grabado -> Cloudinary -> DB attachment -> Deepgram -> GPT -> documento adjunto
    """

    @staticmethod
    async def _download_bytes(url: str) -> bytes:
        async with httpx.AsyncClient(timeout=60.0) as client:
            response = await client.get(url)
            response.raise_for_status()
            return response.content

    @staticmethod
    async def _generate_summary(transcript: str) -> str:
        api_key = settings.OPENAI_API_KEY
        if not api_key:
            return (
                "Se registró la videollamada correctamente, pero no hay OpenAI configurado para generar "
                "el resumen automáticamente."
            )

        client = AsyncOpenAI(api_key=api_key)
        prompt = (
            "Eres un asistente clínico. Genera un resumen breve, claro y útil de esta conversación médica "
            "capturada en una videollamada. Debe incluir: motivo de la consulta, puntos clave discutidos, "
            "preguntas del paciente, decisiones/plan y observaciones importantes. Responde en español.\n\n"
            f"TRANSCRIPCIÓN:\n{transcript[:12000]}"
        )

        completion = await client.chat.completions.create(
            model=settings.OPENAI_MODEL,
            messages=[
                {"role": "system", "content": "Eres un asistente médico que produce resúmenes claros y útiles para documentación clínica."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.2,
            max_tokens=1200,
        )

        return completion.choices[0].message.content.strip() or "Resumen no disponible."

    @staticmethod
    def _create_document_payload(summary: str, transcript: str) -> tuple[bytes, str, str]:
        text = (
            "RESUMEN DE VIDEOLLAMADA\n\n"
            f"{summary}\n\n"
            "--- TRANSCRIPCIÓN COMPLETA ---\n"
            f"{transcript}"
        )

        try:
            from docx import Document

            document = Document()
            document.add_heading("Resumen de videollamada médica", level=0)
            document.add_paragraph(summary)
            document.add_paragraph("\n--- TRANSCRIPCIÓN COMPLETA ---\n")
            document.add_paragraph(transcript)

            buffer = io.BytesIO()
            document.save(buffer)
            buffer.seek(0)
            return buffer.getvalue(), "session_summary.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:
            return text.encode("utf-8"), "session_summary.txt", "text/plain"

    @staticmethod
    async def _upload_to_cloudinary(file_bytes: bytes, filename: str, folder: str, resource_type: str = "video") -> dict:
        stem, extension = os.path.splitext(filename)
        unique_name = f"{stem}_{uuid.uuid4().hex[:10]}"
        public_id = f"{unique_name}{extension}" if resource_type == "raw" else unique_name
        result = cloudinary.uploader.upload(
            io.BytesIO(file_bytes),
            folder=folder,
            resource_type=resource_type,
            public_id=public_id,
        )

        return {
            "public_id": result.get("public_id"),
            "file_url": result.get("secure_url"),
            "secure_url": result.get("secure_url"),
            "file_name": filename,
            "mime_type": result.get("resource_type", "application/octet-stream"),
            "file_size": result.get("bytes", 0),
            "storage_disk": "cloudinary",
        }

    @staticmethod
    async def process_call_recording(
        db: Session,
        conversation_id: int,
        sender_id: int,
        audio_url: str | None = None,
        audio_file: UploadFile | None = None,
        file_name: str = "recording.webm",
    ):
        conversation = db.query(MedicalConversation).filter(
            MedicalConversation.id == conversation_id,
            (
                (MedicalConversation.patient_id == sender_id)
                | (MedicalConversation.doctor_id == sender_id)
            ),
        ).first()

        if not conversation:
            raise ValueError("Conversación no encontrada o no autorizada para este usuario.")

        if audio_file is not None:
            await audio_file.seek(0)
            audio_bytes = await audio_file.read()
            original_filename = audio_file.filename or file_name
            content_type = audio_file.content_type or "audio/webm"
        elif audio_url:
            audio_bytes = await CallTranscriptionPipeline._download_bytes(audio_url)
            original_filename = file_name or "recording.webm"
            content_type = "audio/webm"
        else:
            raise ValueError("Se requiere un audio de la videollamada o una URL con el archivo grabado.")

        folder_path = f"medical_conversations/{conversation_id}/recordings"
        audio_upload = await CallTranscriptionPipeline._upload_to_cloudinary(
            file_bytes=audio_bytes,
            filename=original_filename,
            folder=folder_path,
            resource_type="video",
        )

        message = MedicalConversationService.send_message(
            db=db,
            conversation_id=conversation_id,
            sender_id=sender_id,
            message="[Grabación de videollamada]",
        )

        if not message:
            raise ValueError("No se pudo registrar el mensaje de la videollamada.")

        attachment = MedicalConversationService.create_attachment(
            db=db,
            message_id=message.id,
            attachment_data={
                "file_name": original_filename,
                "file_path": audio_upload["public_id"],
                "file_url": audio_upload["secure_url"],
                "mime_type": content_type,
                "file_size": audio_upload["file_size"],
                "storage_disk": "cloudinary",
                "attachment_type": "audio",
            },
            duration=None,
        )

        transcript = await transcribe_audio(
            audio_bytes,
            content_type=content_type.split(";", 1)[0],
        )
        if not transcript:
            logger.warning(
                "Call recording saved but Deepgram found no speech (conversation_id=%s, audio_bytes=%s)",
                conversation_id,
                len(audio_bytes),
            )
            return {
                "status": "transcription_failed",
                "error": "El audio se guardó en el chat, pero Deepgram no detectó voz. Comprueba que se compartió el audio de la pestaña y vuelve a grabar.",
                "audio_attachment_id": attachment.id,
                "audio_url": attachment.file_url,
                "message_id": message.id,
            }

        summary = await CallTranscriptionPipeline._generate_summary(transcript)

        document_bytes, document_name, document_mime = CallTranscriptionPipeline._create_document_payload(summary, transcript)
        document_upload = await CallTranscriptionPipeline._upload_to_cloudinary(
            file_bytes=document_bytes,
            filename=document_name,
            folder=f"medical_conversations/{conversation_id}/documents",
            resource_type="raw",
        )

        document_message = MedicalConversationService.send_message(
            db=db,
            conversation_id=conversation_id,
            sender_id=sender_id,
            message="[Resumen de videollamada médica]",
        )

        MedicalConversationService.create_attachment(
            db=db,
            message_id=document_message.id,
            attachment_data={
                "file_name": document_name,
                "file_path": document_upload["public_id"],
                "file_url": document_upload["secure_url"],
                "mime_type": document_mime,
                "file_size": document_upload["file_size"],
                "storage_disk": "cloudinary",
                "attachment_type": "document",
            },
            duration=None,
        )

        db.commit()

        return {
            "audio_attachment_id": attachment.id,
            "audio_url": attachment.file_url,
            "transcript": transcript,
            "summary": summary,
            "document_message_id": document_message.id,
            "document_url": document_upload["secure_url"],
            "document_name": document_name,
            "message_id": message.id,
        }